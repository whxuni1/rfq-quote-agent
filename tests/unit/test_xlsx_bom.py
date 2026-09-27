from __future__ import annotations

from collections.abc import Sequence
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment

from rfq_agent.schemas.bom import BOMLine, MaterialClass
from rfq_agent.tools.xlsx_bom import (
    BOMParseError,
    classify_material,
    collect_part_quantities,
    detect_level_strategy,
    is_auto_grade,
    parse_bom,
    single_source_risk,
)

FIXTURES = Path(__file__).parents[1] / "fixtures"


def _xlsx(tmp_path: Path, sheets: dict[str, list[list[object]]], name: str = "bom.xlsx") -> str:
    wb = Workbook()
    first = True
    for title, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        assert ws is not None
        ws.title = title
        first = False
        for row in rows:
            ws.append(row)
    path = tmp_path / name
    wb.save(path)
    return str(path)


def _levels(tree_nodes: Sequence[BOMLine]) -> list[tuple[str, int, str | None]]:
    return [(n.part_no, n.level, n.parent_line_no) for n in tree_nodes]


# ---------------------------------------------------------------------------
# 单 sheet / 多 sheet
# ---------------------------------------------------------------------------


def test_single_sheet_level_column(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "BOM": [
                ["Item", "Level", "Part No", "Description", "Qty", "UoM"],
                [1, 0, "TOP", "Top PCBA assembly", 1, "EA"],
                [2, 1, "MCU-1", "Microcontroller", 1, "ea"],
                [3, 1, "CON-1", "Board connector", 2, None],
            ]
        },
    )
    tree = parse_bom(path, "RFQ-1")
    assert tree.root.part_no == "TOP"
    assert _levels(tree.nodes) == [("TOP", 0, None), ("MCU-1", 1, "1"), ("CON-1", 1, "1")]
    assert tree.nodes[1].uom == "EA"  # 统一大写
    assert tree.nodes[2].uom == "EA"  # 空单位取默认
    assert tree.total_lines == 3
    assert detect_level_strategy(path) == "level_column"


def test_multi_sheet_skips_non_bom_and_header_offset(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "Cover": [["RFQ", "X"], ["Project", "Y"]],
            "BOM": [
                ["Title row"],
                [],
                ["Part No", "Qty", "Description"],
                ["A-1", 1, "Housing"],
            ],
            "Alt": [["Part No", "Qty"], ["A-1", 3]],
        },
    )
    tree = parse_bom(path, "RFQ-2", project_name="Proj")
    # 单行无层级 → 退化策略：合成 root
    assert tree.root.part_desc == "Proj"
    assert tree.root.attributes == {"synthesized": "true"}
    assert [n.part_no for n in tree.nodes] == ["RFQ-2", "A-1"]
    assert collect_part_quantities(path) == {"A-1": [("BOM", Decimal(1)), ("Alt", Decimal(3))]}


def test_no_bom_sheet_raises(tmp_path: Path) -> None:
    path = _xlsx(tmp_path, {"Cover": [["Name", "Value"], ["a", 1]]})
    with pytest.raises(BOMParseError, match="未找到 BOM sheet"):
        parse_bom(path, "R")
    with pytest.raises(BOMParseError):
        detect_level_strategy(path)


def test_bom_sheet_without_rows_raises(tmp_path: Path) -> None:
    path = _xlsx(tmp_path, {"BOM": [["Part No", "Qty"]]})
    with pytest.raises(BOMParseError, match="无数据行"):
        parse_bom(path, "R")


# ---------------------------------------------------------------------------
# 层级策略 1-4
# ---------------------------------------------------------------------------


def test_strategy1_one_based_levels_synthesize_root(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "BOM": [
                ["Level", "Part No", "Qty"],
                [1, "SUB-A", 1],
                [2, "A-1", 2],
                [1, "SUB-B", 1],
            ]
        },
    )
    tree = parse_bom(path, "R")
    assert _levels(tree.nodes) == [
        ("R", 0, None),
        ("SUB-A", 1, "0"),
        ("A-1", 2, "1"),
        ("SUB-B", 1, "0"),
    ]


def test_strategy1_multiple_level0_shifted_under_root(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "BOM": [
                ["Level", "Part No", "Qty"],
                [0, "A", 1],
                [0, "B", 1],
            ]
        },
    )
    tree = parse_bom(path, "R")
    assert _levels(tree.nodes) == [("R", 0, None), ("A", 1, "0"), ("B", 1, "0")]


def test_strategy1_level_jump_raises(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "BOM": [
                ["Level", "Part No", "Qty"],
                [0, "TOP", 1],
                [2, "DEEP", 1],
            ]
        },
    )
    with pytest.raises(BOMParseError, match="层级跳变"):
        parse_bom(path, "R")


def test_strategy1_bad_level_value_raises(tmp_path: Path) -> None:
    path = _xlsx(tmp_path, {"BOM": [["Level", "Part No", "Qty"], ["x", "TOP", 1]]})
    with pytest.raises(BOMParseError, match="层级不是整数"):
        parse_bom(path, "R")


def test_strategy1_missing_level_value_raises(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "BOM": [
                ["Level", "Part No", "Qty"],
                [0, "TOP", 1],
                [None, "A", 1],
            ]
        },
    )
    with pytest.raises(BOMParseError, match="level 列存在空值"):
        parse_bom(path, "R")


def test_strategy2_indent_by_leading_spaces(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "BOM": [
                ["Part No", "Description", "Qty"],
                ["TOP", "Assembly", 1],
                ["  SUB", "Sub assembly", 1],
                ["    LEAF", "Leaf", 4],
                ["  SUB2", "Second", 1],
            ]
        },
    )
    assert detect_level_strategy(path) == "indent"
    tree = parse_bom(path, "R")
    assert _levels(tree.nodes) == [
        ("TOP", 0, None),
        ("SUB", 1, "1"),
        ("LEAF", 2, "2"),
        ("SUB2", 1, "1"),
    ]


def test_strategy2_indent_by_cell_alignment(tmp_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(["Part No", "Qty"])
    ws.append(["TOP", 1])
    ws.append(["CHILD", 1])
    ws["A3"].alignment = Alignment(indent=1)
    path = tmp_path / "indent.xlsx"
    wb.save(path)
    tree = parse_bom(str(path), "R")
    assert _levels(tree.nodes) == [("TOP", 0, None), ("CHILD", 1, "1")]


def test_strategy3_parent_column(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "BOM": [
                ["Line", "Part No", "Parent", "Qty"],
                ["10", "TOP", None, 1],
                ["20", "SUB", "TOP", 1],
                ["30", "LEAF", "SUB", 2],
            ]
        },
    )
    assert detect_level_strategy(path) == "parent_column"
    tree = parse_bom(path, "R")
    assert _levels(tree.nodes) == [("TOP", 0, None), ("SUB", 1, "10"), ("LEAF", 2, "20")]


def test_strategy3_multiple_tops_synthesize_root(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "BOM": [
                ["Part No", "Parent", "Qty"],
                ["A", None, 1],
                ["B", None, 1],
                ["A1", "A", 1],
            ]
        },
    )
    tree = parse_bom(path, "R")
    assert _levels(tree.nodes) == [
        ("R", 0, None),
        ("A", 1, "0"),
        ("B", 1, "0"),
        ("A1", 2, "1"),
    ]


def test_strategy3_unknown_parent_raises(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "BOM": [
                ["Part No", "Parent", "Qty"],
                ["TOP", None, 1],
                ["A", "NOPE", 1],
            ]
        },
    )
    with pytest.raises(BOMParseError, match="父件号不存在"):
        parse_bom(path, "R")


def test_strategy3_cycle_raises(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "BOM": [
                ["Part No", "Parent", "Qty"],
                ["TOP", None, 1],
                ["A", "B", 1],
                ["B", "A", 1],
            ]
        },
    )
    with pytest.raises(BOMParseError, match="成环"):
        parse_bom(path, "R")


def test_strategy4_flat_fallback(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "BOM": [
                ["Line", "Part No", "Qty"],
                ["0", "A", 1],
                ["ROOT", "B", 1],
            ]
        },
    )
    assert detect_level_strategy(path) == "flat"
    tree = parse_bom(path, "RFQ-9")
    # "0"、"ROOT" 已被占用 → root 行号退到下一个候选
    assert tree.root.line_no == "ROOT-0"
    assert tree.root.part_desc == "RFQ-9"
    assert all(n.level == 1 and n.parent_line_no == "ROOT-0" for n in tree.nodes[1:])


# ---------------------------------------------------------------------------
# 异常列名 / 数据错误
# ---------------------------------------------------------------------------


def test_chinese_and_irregular_headers(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "物料清单": [
                [
                    "序号",
                    "层级",
                    "物料号.",
                    " 物料描述 ",
                    "用量",
                    "单位",
                    "指定供应商",
                    "规格",
                    "图号",
                    "币种",
                    "Target Price (CNY)",
                ],
                ["1", 0, "TOP", "域控制器总成", 1, "EA", None, None, None, "CNY", None],
                [
                    "2",
                    1,
                    "SOC-1",
                    "SoC 芯片 车规",
                    1,
                    "EA",
                    "Qualcomm",
                    "process_node=5nm; cores=8",
                    "DWG-9",
                    None,
                    100,
                ],
                ["3", 1, "X-1", "Misc part", 1, "EA", "Open", "free text spec", None, None, None],
            ]
        },
    )
    tree = parse_bom(path, "R")
    soc = tree.nodes[1]
    assert soc.material_class is MaterialClass.SOC
    assert soc.auto_grade_required is True
    assert soc.single_source_risk is True
    assert soc.attributes == {
        "process_node": "5nm",
        "cores": "8",
        "drawing": "DWG-9",
        "Target Price (CNY)": "100",
    }
    assert tree.nodes[2].attributes == {"spec": "free text spec"}
    assert tree.currency_hint == "CNY"


def test_currency_from_header_suffix(tmp_path: Path) -> None:
    path = _xlsx(
        tmp_path,
        {
            "BOM": [
                ["Part No", "Qty", "Unit Price (EUR)"],
                ["A", 1, 2.5],
            ]
        },
    )
    assert parse_bom(path, "R").currency_hint == "EUR"


@pytest.mark.parametrize("qty, message", [("abc", "不是数字"), (0, "必须 > 0"), (-1, "必须 > 0")])
def test_bad_qty_raises_with_location(tmp_path: Path, qty: object, message: str) -> None:
    path = _xlsx(tmp_path, {"BOM": [["Part No", "Qty"], ["A", qty]]})
    with pytest.raises(BOMParseError, match=message) as info:
        parse_bom(path, "R")
    assert info.value.location == "BOM!R2"


def test_missing_part_no_raises(tmp_path: Path) -> None:
    path = _xlsx(tmp_path, {"BOM": [["Part No", "Qty", "Desc"], [None, 1, "orphan"]]})
    with pytest.raises(BOMParseError, match="缺少 Part No"):
        parse_bom(path, "R")


def test_duplicate_line_no_raises(tmp_path: Path) -> None:
    path = _xlsx(tmp_path, {"BOM": [["Line", "Part No", "Qty"], ["1", "A", 1], ["1", "B", 1]]})
    with pytest.raises(BOMParseError, match="行号重复"):
        parse_bom(path, "R")


def test_missing_and_corrupt_file(tmp_path: Path) -> None:
    with pytest.raises(BOMParseError, match="文件不存在"):
        parse_bom(str(tmp_path / "nope.xlsx"), "R")
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(BOMParseError, match="无法打开"):
        parse_bom(str(bad), "R")


# ---------------------------------------------------------------------------
# 规则函数
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "text, expected",
    [
        ("Cockpit SoC application processor", MaterialClass.SOC),
        ("Socket adapter", MaterialClass.OTHER),  # "soc" 词边界
        ("LPDDR5 16GB DRAM", MaterialClass.MEMORY),
        ("DDR4 8GB", MaterialClass.MEMORY),
        ("Camera housing", MaterialClass.HOUSING),
        ("Display cover glass", MaterialClass.OPTICS),
        ("8MP CMOS image sensor", MaterialClass.CAMERA),
        ("Radar sensor 77GHz", MaterialClass.SENSOR),
        ("Main PCBA", MaterialClass.PCBA),
        ("Lenses set", MaterialClass.OPTICS),
        ("2 Optical Requirements", MaterialClass.OPTICS),
        ("散热片", MaterialClass.THERMAL),
        ("Android software license", MaterialClass.SOFTWARE),
        ("", MaterialClass.OTHER),
    ],
)
def test_classify_material(text: str, expected: MaterialClass) -> None:
    assert classify_material(text) is expected


def test_classify_material_custom_rules() -> None:
    rules = ((MaterialClass.MCU, ("widget",)),)
    assert classify_material("blue widget", rules) is MaterialClass.MCU
    assert classify_material("SoC", rules) is MaterialClass.OTHER


@pytest.mark.parametrize(
    "hint, expected",
    [
        (None, None),
        ("  ", None),
        ("Qualcomm", True),
        ("Open", False),
        ("TBD", False),
        ("Micron / Samsung", False),
        ("TE, Molex", False),
        ("Sony or Onsemi", False),
    ],
)
def test_single_source_risk(hint: str | None, expected: bool | None) -> None:
    assert single_source_risk(hint) is expected


def test_is_auto_grade() -> None:
    assert is_auto_grade("grade=AEC-Q100")
    assert is_auto_grade("车规级")
    assert not is_auto_grade("consumer grade")


# ---------------------------------------------------------------------------
# fixture / 确定性
# ---------------------------------------------------------------------------


def test_case2_fixture_three_levels() -> None:
    tree = parse_bom(str(FIXTURES / "case2_cockpit_core" / "bom.xlsx"), "case2")
    assert tree.root.part_no == "CDC-X1-ASSY"
    assert max(n.level for n in tree.nodes) >= 2
    level2 = {n.material_class for n in tree.nodes if n.level == 2}
    assert {MaterialClass.SOC, MaterialClass.MEMORY} <= level2
    level2_desc = " ".join(n.part_desc for n in tree.nodes if n.level == 2)
    assert "LPDDR5" in level2_desc and "NAND" in level2_desc


def test_parse_is_deterministic() -> None:
    path = str(FIXTURES / "case3_adas_camera" / "bom.xlsx")
    assert parse_bom(path, "case3") == parse_bom(path, "case3")
