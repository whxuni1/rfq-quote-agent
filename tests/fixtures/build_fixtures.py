"""生成 3 个 RFQ fixture（spec.pdf + bom.xlsx + 图纸占位 + LLM 回放）。

用法：python tests/fixtures/build_fixtures.py
产物已提交到 git；修改本脚本后需重新生成并一并提交。
PDF 用内置极简写入器生成（Helvetica，文本可抽取），不引入额外第三方库。
"""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment

ROOT = Path(__file__).parent

# ---------------------------------------------------------------------------
# 极简 PDF 写入器
# ---------------------------------------------------------------------------


class MiniPDF:
    WIDTH, HEIGHT, MARGIN = 595, 842, 50

    def __init__(self) -> None:
        self.pages: list[list[str]] = [[]]
        self.y: float = self.HEIGHT - self.MARGIN

    def _ensure(self, height: float) -> None:
        if self.y - height < self.MARGIN:
            self.pages.append([])
            self.y = self.HEIGHT - self.MARGIN

    @staticmethod
    def _esc(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    def _text(self, x: float, y: float, text: str, size: float, bold: bool = False) -> None:
        font = "F2" if bold else "F1"
        self.pages[-1].append(f"BT /{font} {size} Tf {x:.1f} {y:.1f} Td ({self._esc(text)}) Tj ET")

    def heading(self, text: str, size: float = 14) -> None:
        self._ensure(size * 2.2)
        self.y -= size * 1.6
        self._text(self.MARGIN, self.y, text, size, bold=True)
        self.y -= size * 0.6

    def para(self, text: str, size: float = 10) -> None:
        self._ensure(size * 1.5)
        self.y -= size * 1.5
        self._text(self.MARGIN, self.y, text, size)

    def table(self, rows: list[list[str]], widths: list[float], size: float = 9) -> None:
        row_h = size * 2
        self._ensure(row_h * len(rows) + 10)
        self.y -= 6
        top = self.y
        x_edges = [float(self.MARGIN)]
        for w in widths:
            x_edges.append(x_edges[-1] + w)
        ops = self.pages[-1]
        for i in range(len(rows) + 1):
            yy = top - i * row_h
            ops.append(f"{x_edges[0]:.1f} {yy:.1f} m {x_edges[-1]:.1f} {yy:.1f} l S")
        bottom = top - len(rows) * row_h
        for xx in x_edges:
            ops.append(f"{xx:.1f} {top:.1f} m {xx:.1f} {bottom:.1f} l S")
        for r, row in enumerate(rows):
            for c, cell in enumerate(row):
                self._text(x_edges[c] + 4, top - (r + 1) * row_h + size * 0.7, cell, size)
        self.y = bottom - 4

    def save(self, path: Path) -> None:
        objects: list[bytes] = []
        n_pages = len(self.pages)
        page_ids = [5 + 2 * i for i in range(n_pages)]
        objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
        kids = " ".join(f"{pid} 0 R" for pid in page_ids)
        objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode())
        objects.append(
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
        )
        objects.append(
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold "
            b"/Encoding /WinAnsiEncoding >>"
        )
        for i, ops in enumerate(self.pages):
            content = "\n".join(ops).encode("latin-1")
            objects.append(
                f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 {self.WIDTH} {self.HEIGHT}] "
                f"/Resources << /Font << /F1 3 0 R /F2 4 0 R >> >> "
                f"/Contents {page_ids[i] + 1} 0 R >>".encode()
            )
            objects.append(
                f"<< /Length {len(content)} >>\nstream\n".encode() + content + b"\nendstream"
            )
        out = bytearray(b"%PDF-1.4\n")
        offsets = []
        for num, body in enumerate(objects, start=1):
            offsets.append(len(out))
            out += f"{num} 0 obj\n".encode() + body + b"\nendobj\n"
        xref = len(out)
        out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
        for off in offsets:
            out += f"{off:010d} 00000 n \n".encode()
        trailer = f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
        out += trailer.encode()
        path.write_bytes(bytes(out))


def _common_tail(pdf: MiniPDF, payment: str, warranty: str) -> None:
    pdf.heading("6 Quality Requirements")
    pdf.para("Supplier shall be certified to IATF 16949 and ISO 9001.")
    pdf.para("Submission: PPAP Level 3 with full APQP deliverables.")
    pdf.heading("7 Commercial Terms")
    pdf.para(f"- Payment terms: {payment}.")
    pdf.para(f"- Warranty: {warranty}.")
    pdf.para("- Quotation validity: 90 days.")
    pdf.para("- Tooling shall be owned by customer after full payment.")


# ---------------------------------------------------------------------------
# case1：简单 PCBA + 结构件（level 列，两层，无歧义）
# ---------------------------------------------------------------------------


def build_case1(d: Path) -> None:
    pdf = MiniPDF()
    pdf.heading("RFQ TECHNICAL SPECIFICATION", size=18)
    pdf.para("Customer: Acme Automotive Co.")
    pdf.para("Project: BCM Gen2 Body Control Module")
    pdf.para("RFQ Reference: RFQ-2026-001")
    pdf.heading("1 Scope")
    pdf.para("Supply of the body control module PCBA with plastic housing, including")
    pdf.para("final assembly, end-of-line test and packaging. Reference drawing DWG-1001 Rev A.")
    pdf.heading("2 Housing")
    pdf.para("PA66-GF30 housing, IP5K4, per drawing DWG-1001.")
    pdf.heading("3 Quantity")
    pdf.para("Annual volume: 200,000 pcs. Project life: 5 years.")
    pdf.heading("4 Delivery")
    pdf.para("SOP 2027-06. Incoterm FCA supplier plant.")
    pdf.para("Delivery location: Wuhu, China")
    pdf.heading("5 Packaging")
    pdf.para("Returnable ESD trays, 40 pcs per tray.")
    _common_tail(pdf, payment="60 days net", warranty="24 months from SOP")
    pdf.save(d / "spec.pdf")

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "BOM"
    ws.append(
        [
            "Item",
            "Level",
            "Part No",
            "Description",
            "Qty",
            "UoM",
            "Vendor",
            "Drawing",
            "Attributes",
            "Currency",
        ]
    )
    rows: list[list[Any]] = [
        [
            "1",
            0,
            "BCM-ASSY-01",
            "BCM Gen2 main PCBA assembly",
            1,
            "EA",
            None,
            "DWG-1001",
            None,
            "USD",
        ],
        ["2", 1, "PCB-4L-120", "4-layer PCB 120x80mm", 1, "EA", "Open", None, "layers=4", None],
        [
            "3",
            1,
            "S32K144",
            "32-bit MCU automotive grade AEC-Q100",
            1,
            "EA",
            "NXP / Infineon",
            None,
            "flash_kb=512",
            None,
        ],
        ["4", 1, "CON-26P", "26-pin board connector", 2, "EA", "TE, Molex", None, None, None],
        ["5", 1, "HSG-PA66", "PA66-GF30 housing", 1, "EA", "Open", "DWG-1001", None, None],
        ["6", 1, "PKG-TRAY", "ESD packaging tray", 0.025, "EA", "Open", None, None, None],
    ]
    for row in rows:
        ws.append(row)
    wb.save(d / "bom.xlsx")

    (d / "drawings").mkdir(exist_ok=True)
    (d / "drawings" / "DWG-1001_RevA.pdf").write_bytes(b"%PDF-1.4\n% drawing placeholder\n")

    _write_llm(
        d,
        {
            "quote_intent": {
                "project_type": "NPI+Mass",
                "annual_volume": "200000",
                "peak_monthly_volume": "20000",
                "project_life_years": 5,
                "currency": "USD",
                "incoterm_target": "FCA",
                "payment_terms_target": "60 days net",
                "nre_scope": ["tooling", "validation"],
                "special_reqs": ["24 months warranty"],
                "skip_agents": [],
            },
            "ambiguity_flags": [],
        },
    )


# ---------------------------------------------------------------------------
# case2：智能座舱核心板（缩进层级，3 层，多 sheet 用量冲突）
# ---------------------------------------------------------------------------


def build_case2(d: Path) -> None:
    pdf = MiniPDF()
    pdf.heading("RFQ TECHNICAL SPECIFICATION", size=18)
    pdf.para("Customer: Nova EV Motors")
    pdf.para("Project: CDC-X1 Cockpit Domain Controller")
    pdf.heading("1 Scope")
    pdf.para("Cockpit domain controller driving 2 displays: core board, carrier board,")
    pdf.para("die-cast housing, heatsink and harness. All ICs shall be AEC-Q100 qualified.")
    pdf.heading("2 Thermal Requirements")
    pdf.para("Operating temperature -40 to 85 C ambient; SoC junction below 105 C.")
    pdf.heading("3 Quantity")
    pdf.table(
        [
            ["Item", "Value"],
            ["Annual volume", "150,000"],
            ["Peak monthly", "15,000"],
            ["Project life", "6 years"],
        ],
        widths=[160, 160],
    )
    pdf.heading("4 Delivery")
    pdf.para("SOP Q3 2027. Incoterm DAP customer plant.")
    pdf.para("Delivery location: Hefei, China")
    pdf.heading("5 Software")
    pdf.para("Hypervisor and Android Automotive license per unit, quoted separately.")
    _common_tail(pdf, payment="90 days after invoice", warranty="36 months or 100,000 km")
    pdf.save(d / "spec.pdf")

    wb = Workbook()
    cover = wb.active
    assert cover is not None
    cover.title = "Cover"
    cover.append(["RFQ", "RFQ-2026-002"])
    cover.append(["Project", "CDC-X1"])

    bom = wb.create_sheet("BOM")
    bom.append(["BOM for CDC-X1 (USD)"])
    bom.append([])
    bom.append(["Part No.", "Description", "Qty", "Unit", "Source", "Spec", "Target Price (USD)"])
    rows: list[tuple[int, list[Any]]] = [
        (0, ["CDC-X1-ASSY", "Cockpit domain controller assembly", 1, "EA", None, None, None]),
        (1, ["CORE-BRD-01", "Core board PCBA 10-layer HDI", 1, "EA", None, None, None]),
        (
            2,
            [
                "SA8295P",
                "Cockpit SoC application processor",
                1,
                "EA",
                "Qualcomm",
                "process_node=5nm; cores=8; grade=AEC-Q100",
                180,
            ],
        ),
        (
            2,
            [
                "LP5-16G",
                "LPDDR5 16GB DRAM",
                1,
                "EA",
                "Micron / Samsung",
                "type=LPDDR5; capacity_gb=16; speed_mbps=6400",
                38,
            ],
        ),
        (
            2,
            [
                "UFS-256",
                "UFS 3.1 256GB NAND storage",
                1,
                "EA",
                "Kioxia, Samsung, SK Hynix",
                "capacity_gb=256",
                22,
            ],
        ),
        (2, ["PMIC-8295", "Power management IC", 2, "EA", "Qualcomm", None, 4]),
        (1, ["CARRIER-01", "Carrier board PCBA 6-layer", 1, "EA", None, None, None]),
        (2, ["SER-GMSL", "GMSL video serializer", 2, "EA", "Analog Devices", None, 6]),
        (2, ["CON-FAKRA", "FAKRA connector quad", 3, "EA", "Rosenberger / TE", None, 3]),
        (1, ["HSG-ADC12", "Die-cast aluminium housing", 1, "EA", "Open", None, 14]),
        (1, ["HS-01", "Heatsink with thermal pad", 1, "EA", "Open", None, 5]),
        (1, ["HRN-PWR", "Power harness 0.6m", 1, "EA", "Yazaki / Sumitomo", None, 2]),
        (1, ["SW-AAOS", "Android Automotive software license", 1, "EA", "Google", None, 9]),
    ]
    for level, values in rows:
        values[0] = "  " * level + str(values[0])
        bom.append(values)
    for row in bom.iter_rows(min_row=4):
        row[1].alignment = Alignment(horizontal="left")

    options = wb.create_sheet("Memory Options")
    options.append(["Part No", "Description", "Qty"])
    options.append(["LP5-16G", "LPDDR5 16GB DRAM (2x8GB variant)", 2])
    wb.save(d / "bom.xlsx")

    _write_llm(
        d,
        {
            "quote_intent": {
                "project_type": "NPI+Mass",
                "annual_volume": "150000",
                "peak_monthly_volume": "15000",
                "project_life_years": 6,
                "currency": "USD",
                "incoterm_target": "DAP",
                "payment_terms_target": "90 days after invoice",
                "nre_scope": ["tooling", "validation", "software_license"],
                "special_reqs": ["36 months or 100,000 km warranty", "all ICs AEC-Q100"],
                "skip_agents": [],
            },
            "ambiguity_flags": [
                {
                    "flag_id": "LLM-1",
                    "category": "term_undefined",
                    "description": "Software license fee basis (per unit vs lump sum) not defined",
                    "related_line_no": None,
                    "severity": "medium",
                },
            ],
        },
    )


# ---------------------------------------------------------------------------
# case3：ADAS 前视摄像头模组（父件列层级，缺光学物料，图纸缺失）
# ---------------------------------------------------------------------------


def build_case3(d: Path) -> None:
    pdf = MiniPDF()
    pdf.heading("RFQ TECHNICAL SPECIFICATION", size=18)
    pdf.para("Customer: Orion Mobility GmbH")
    pdf.para("Project: FCM-8 ADAS Front Camera Module")
    pdf.heading("1 Scope")
    pdf.para("8MP front camera module with image sensor, serializer, lens assembly and")
    pdf.para("housing. Outline per drawing DWG-3001 Rev B, lens mount per DWG-3002.")
    pdf.heading("2 Optical Requirements")
    pdf.para("Lens: 8MP, HFOV 120 deg, F1.8, IR-cut filter, active alignment required.")
    pdf.heading("3 Quantity")
    pdf.para("Annual volume: 300k units. Project life: 7 years.")
    pdf.heading("4 Delivery")
    pdf.para("SOP 2028-01. Incoterm DDP customer plant.")
    pdf.para("Delivery location: Ingolstadt, Germany")
    pdf.para("All prices shall be quoted in EUR.")
    _common_tail(pdf, payment="45 days net", warranty="60 months")
    pdf.save(d / "spec.pdf")

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Camera BOM"
    ws.append(["Line", "Part No", "Parent", "Description", "Qty", "UoM", "Supplier"])
    rows: list[list[Any]] = [
        ["10", "FCM-8-ASSY", None, "Front camera module assembly", 1, "EA", None],
        ["20", "CAM-PCBA-01", "FCM-8-ASSY", "Camera PCBA 4-layer", 1, "EA", None],
        ["30", "IMX728", "CAM-PCBA-01", "8MP CMOS image sensor AEC-Q100", 1, "EA", "Sony"],
        ["40", "TBD", "CAM-PCBA-01", "GMSL2 serializer", 1, "EA", "Analog Devices"],
        ["50", "HSG-CAM", "FCM-8-ASSY", "Aluminium camera housing", 1, "EA", "Open"],
        ["60", "CON-FAKRA-1", "FCM-8-ASSY", "FAKRA connector", 1, "M", "Rosenberger"],
    ]
    for row in rows:
        ws.append(row)
    wb.save(d / "bom.xlsx")

    (d / "drawings").mkdir(exist_ok=True)
    (d / "drawings" / "DWG-3001_RevB.pdf").write_bytes(b"%PDF-1.4\n% drawing placeholder\n")

    _write_llm(
        d,
        {
            "quote_intent": {
                "project_type": "NPI+Mass",
                "annual_volume": "300000",
                "peak_monthly_volume": "30000",
                "project_life_years": 7,
                "currency": "EUR",
                "incoterm_target": "DDP",
                "payment_terms_target": "45 days net",
                "nre_scope": ["tooling", "jigs", "validation"],
                "special_reqs": ["60 months warranty", "active alignment"],
                "skip_agents": [],
            },
            "ambiguity_flags": [
                {
                    "flag_id": "LLM-1",
                    "category": "spec_gap",
                    "description": (
                        "Spec requires IR-cut filter and 120 deg lens "
                        "but BOM has no lens or filter line"
                    ),
                    "related_line_no": None,
                    "severity": "high",
                },
            ],
        },
    )


def _write_llm(d: Path, response: dict[str, object]) -> None:
    (d / "llm_ingestion.json").write_text(
        json.dumps([response], ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )


def main() -> None:
    for name, builder in (
        ("case1_simple_pcba", build_case1),
        ("case2_cockpit_core", build_case2),
        ("case3_adas_camera", build_case3),
    ):
        d = ROOT / name
        d.mkdir(exist_ok=True)
        builder(d)
        print(f"built {d}")


if __name__ == "__main__":
    main()
