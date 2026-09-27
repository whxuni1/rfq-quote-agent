"""Excel BOM → BOMTree。纯结构解析，零语义理解。

层级识别策略（按优先级）：
1. 显式 level 列（"Level"/"层级"）
2. 缩进（part no / 描述单元格的前导空白或单元格缩进格式）
3. 父件号列（"Parent"/"父件"）
4. 退化：全部 level=1，root 用项目名构造
"""

from __future__ import annotations

import re
from collections.abc import Iterable, Sequence
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Literal

import structlog
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pydantic import ValidationError

from rfq_agent.schemas.bom import BOMLine, BOMTree, MaterialClass
from rfq_agent.tools import ToolError

# ---------------------------------------------------------------------------
# 可配置常量
# ---------------------------------------------------------------------------

# 表头同义词（比较前做归一化：小写、去首尾空白、压缩空白、去掉结尾的 . : ：）
HEADER_SYNONYMS: dict[str, tuple[str, ...]] = {
    "line_no": ("item", "line", "line no", "pos", "position", "序号", "行号"),
    "level": ("level", "lvl", "bom level", "层级", "阶层"),
    "parent": ("parent", "parent part", "parent part no", "parent p/n", "父件", "父件号"),
    "part_no": ("part no", "part number", "p/n", "pn", "物料号", "料号", "零件号"),
    "part_desc": ("description", "desc", "part description", "名称", "描述", "物料描述"),
    "qty": ("qty", "quantity", "qty per", "用量", "数量"),
    "uom": ("uom", "unit", "单位"),
    "source": (
        "source",
        "vendor",
        "manufacturer",
        "mfr",
        "supplier",
        "供应商",
        "厂商",
        "指定供应商",
    ),
    "drawing": ("drawing", "drawing no", "dwg", "图纸", "图号"),
    "currency": ("currency", "币种"),
    "attributes": ("attributes", "spec", "specification", "规格", "参数"),
}

# 物料分类规则：按顺序匹配，先命中先得。顺序即优先级（如 "cover glass" 先于 "cover"）。
MATERIAL_CLASS_RULES: tuple[tuple[MaterialClass, tuple[str, ...]], ...] = (
    (MaterialClass.SOFTWARE, ("software", "license", "licence", "firmware", "软件", "许可")),
    (MaterialClass.PACKAGING, ("packaging", "carton", "tray", "blister", "包装")),
    (
        MaterialClass.OPTICS,
        (
            "lens",
            "optic",
            "optical",
            "ir filter",
            "ir-cut",
            "cover glass",
            "镜头",
            "光学",
            "滤光片",
        ),
    ),
    (
        MaterialClass.THERMAL,
        ("heatsink", "heat sink", "thermal", "fan", "vapor chamber", "散热", "导热"),
    ),
    (
        MaterialClass.HOUSING,
        ("housing", "enclosure", "bracket", "cover", "shell", "frame", "外壳", "支架", "壳体"),
    ),
    (MaterialClass.HARNESS, ("harness", "cable assembly", "wire", "线束")),
    (MaterialClass.CONNECTOR, ("connector", "fakra", "hsd", "receptacle", "连接器")),
    (MaterialClass.CAMERA, ("camera", "image sensor", "cmos sensor", "imager", "摄像头")),
    (MaterialClass.DISPLAY, ("display", "lcd", "tft", "oled", "touch panel", "显示", "屏")),
    (MaterialClass.SOC, ("soc", "system on chip", "application processor", "isp", "芯片组")),
    (MaterialClass.MCU, ("mcu", "microcontroller", "单片机")),
    (
        MaterialClass.MEMORY,
        (
            "ddr",
            "lpddr",
            "nand",
            "nor flash",
            "emmc",
            "ufs",
            "dram",
            "flash",
            "memory",
            "内存",
            "存储",
        ),
    ),
    (MaterialClass.SENSOR, ("sensor", "radar", "imu", "gyro", "accelerometer", "传感器")),
    (
        MaterialClass.PCBA,
        ("pcba", "pcb", "circuit board", "main board", "mainboard", "主板", "电路板"),
    ),
)

AUTO_GRADE_KEYWORDS: tuple[str, ...] = (
    "aec-q100",
    "aec-q101",
    "aec-q104",
    "aec-q200",
    "automotive grade",
    "automotive-grade",
    "车规",
)

# source_hint 中表示"不限源"的词，出现即视为非单源
OPEN_SOURCE_TOKENS: frozenset[str] = frozenset({"open", "any", "tbd", "multi", "n/a", "不限"})

HEADER_SCAN_ROWS = 20
DEFAULT_UOM = "EA"
_SOURCE_SPLIT = re.compile(r"[/,;|、，；]|\s+or\s+|\s+and\s+|\s+&\s+", re.IGNORECASE)
_ISO_CCY_IN_HEADER = re.compile(r"\(([A-Z]{3})\)")

LevelStrategy = Literal["level_column", "indent", "parent_column", "flat"]


class BOMParseError(ToolError):
    """BOM 解析失败。`location` 指出 sheet/行号，便于调用方定位。"""

    def __init__(self, message: str, location: str | None = None) -> None:
        super().__init__(f"{message} [{location}]" if location else message)
        self.location = location


# ---------------------------------------------------------------------------
# 公共函数
# ---------------------------------------------------------------------------


def classify_material(
    text: str,
    rules: Sequence[tuple[MaterialClass, Sequence[str]]] = MATERIAL_CLASS_RULES,
) -> MaterialClass:
    """关键词规则分类。ASCII 关键词按词边界匹配（允许后接数字/复数），中文按子串匹配。"""
    lowered = text.lower()
    for material_class, keywords in rules:
        if any(_keyword_hit(lowered, kw) for kw in keywords):
            return material_class
    return MaterialClass.OTHER


def is_auto_grade(text: str, keywords: Sequence[str] = AUTO_GRADE_KEYWORDS) -> bool:
    lowered = text.lower()
    return any(kw in lowered for kw in keywords)


def single_source_risk(source_hint: str | None) -> bool | None:
    """source_hint 指定单一厂商且无备选 → True；多个厂商或不限源 → False；未知 → None。"""
    if source_hint is None or not source_hint.strip():
        return None
    if source_hint.strip().lower() in OPEN_SOURCE_TOKENS:
        return False
    vendors = [p.strip() for p in _SOURCE_SPLIT.split(source_hint) if p.strip()]
    return len(vendors) == 1


def parse_bom(xlsx_path: str, rfq_id: str, project_name: str | None = None) -> BOMTree:
    """解析 Excel BOM。取第一个含 part no + qty 表头的 sheet 作为主 BOM。"""
    log = structlog.get_logger().bind(tool=__name__, rfq_id=rfq_id, step="parse_bom")
    workbook = _open(xlsx_path)
    for ws in workbook.worksheets:
        header = _find_header(ws)
        if header is None:
            continue
        rows = _read_rows(ws, header)
        if not rows:
            raise BOMParseError("BOM sheet 无数据行", location=ws.title)
        strategy = _pick_strategy(header, rows)
        log.debug("bom_sheet_found", sheet=ws.title, rows=len(rows), strategy=strategy)
        tree = _build_tree(rows, strategy, rfq_id, project_name or rfq_id, header.currency)
        log.info("bom_parsed", sheet=ws.title, total_lines=tree.total_lines, strategy=strategy)
        return tree
    raise BOMParseError("未找到 BOM sheet（需同时含 Part No 与 Qty 列）", location=xlsx_path)


def detect_level_strategy(xlsx_path: str) -> LevelStrategy:
    """返回主 BOM sheet 采用的层级识别策略（审计/测试用）。"""
    workbook = _open(xlsx_path)
    for ws in workbook.worksheets:
        header = _find_header(ws)
        if header is not None:
            return _pick_strategy(header, _read_rows(ws, header))
    raise BOMParseError("未找到 BOM sheet", location=xlsx_path)


def collect_part_quantities(xlsx_path: str) -> dict[str, list[tuple[str, Decimal]]]:
    """扫描全部 BOM 类 sheet，返回 part_no → [(sheet, qty)]，供跨 sheet 用量一致性检查。"""
    workbook = _open(xlsx_path)
    result: dict[str, list[tuple[str, Decimal]]] = {}
    for ws in workbook.worksheets:
        header = _find_header(ws)
        if header is None:
            continue
        for row in _read_rows(ws, header):
            result.setdefault(row.part_no, []).append((ws.title, row.qty))
    return result


# ---------------------------------------------------------------------------
# 内部实现
# ---------------------------------------------------------------------------


@dataclass
class _Header:
    row_idx: int  # 1-based
    columns: dict[str, int]  # 字段 → 0-based 列号
    extra: dict[int, str]  # 未识别列 → 原始表头，进 attributes
    currency: str | None


@dataclass
class _Row:
    excel_row: int
    line_no: str | None
    level: int | None
    parent: str | None
    part_no: str
    part_desc: str
    qty: Decimal
    uom: str
    source: str | None
    indent: int
    attributes: dict[str, str] = field(default_factory=dict)


def _open(xlsx_path: str) -> Any:
    if not Path(xlsx_path).is_file():
        raise BOMParseError("文件不存在", location=xlsx_path)
    try:
        return load_workbook(xlsx_path, data_only=True)
    except Exception as exc:  # openpyxl 对坏文件抛多种异常，统一转 ToolError
        raise BOMParseError(f"无法打开 Excel: {exc}", location=xlsx_path) from exc


def _norm_header(value: object) -> str:
    text = re.sub(r"\s+", " ", str(value)).strip().lower()
    return text.rstrip(".:：").strip()


def _keyword_hit(lowered: str, keyword: str) -> bool:
    if keyword.isascii():
        pattern = r"(?<![a-z0-9])" + re.escape(keyword) + r"(?:s|es)?(?![a-z])"
        return re.search(pattern, lowered) is not None
    return keyword in lowered


def _find_header(ws: Worksheet) -> _Header | None:
    lookup = {syn: key for key, syns in HEADER_SYNONYMS.items() for syn in syns}
    for row_idx, values in enumerate(
        ws.iter_rows(min_row=1, max_row=HEADER_SCAN_ROWS, values_only=True), start=1
    ):
        columns: dict[str, int] = {}
        extra: dict[int, str] = {}
        currency: str | None = None
        for col_idx, value in enumerate(values):
            if value is None or not str(value).strip():
                continue
            key = lookup.get(_norm_header(value))
            if key is not None and key not in columns:
                columns[key] = col_idx
            else:
                extra[col_idx] = str(value).strip()
                match = _ISO_CCY_IN_HEADER.search(str(value))
                if match and currency is None:
                    currency = match.group(1)
        if "part_no" in columns and "qty" in columns:
            return _Header(row_idx=row_idx, columns=columns, extra=extra, currency=currency)
    return None


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _leading_indent(raw: str) -> int:
    expanded = raw.expandtabs(4)
    return len(expanded) - len(expanded.lstrip(" "))


def _parse_attributes(raw: str) -> dict[str, str] | None:
    """解析 "k=v; k=v" 或 "k: v, k: v"；无法解析返回 None。"""
    parts = [p.strip() for p in re.split(r"[;；\n]", raw) if p.strip()]
    if len(parts) == 1 and "," in parts[0]:
        parts = [p.strip() for p in parts[0].split(",") if p.strip()]
    result: dict[str, str] = {}
    for part in parts:
        match = re.match(r"^([^=:：]+)[=:：](.+)$", part)
        if match is None:
            return None
        result[match.group(1).strip()] = match.group(2).strip()
    return result or None


def _read_rows(ws: Worksheet, header: _Header) -> list[_Row]:
    cols = header.columns
    rows: list[_Row] = []
    for cells in ws.iter_rows(min_row=header.row_idx + 1):
        values = [c.value for c in cells]
        if all(v is None or not str(v).strip() for v in values):
            continue
        excel_row = int(cells[0].row or 0)
        loc = f"{ws.title}!R{excel_row}"

        def get(key: str, _values: list[Any] = values) -> str:
            idx = cols.get(key)
            return _cell_text(_values[idx]) if idx is not None and idx < len(_values) else ""

        raw_part_no = get("part_no")
        if not raw_part_no.strip():
            raise BOMParseError("缺少 Part No", location=loc)
        raw_desc = get("part_desc")
        indent = _leading_indent(raw_part_no) or _leading_indent(raw_desc)
        if not indent:
            for key in ("part_no", "part_desc"):
                idx = cols.get(key)
                if idx is not None and idx < len(cells):
                    alignment_indent = cells[idx].alignment.indent
                    if alignment_indent:
                        indent = int(alignment_indent)
                        break

        qty_text = get("qty").strip()
        try:
            qty = Decimal(qty_text)
        except InvalidOperation as exc:
            raise BOMParseError(f"用量不是数字: {qty_text!r}", location=loc) from exc
        if not qty.is_finite() or qty <= 0:
            raise BOMParseError(f"用量必须 > 0: {qty_text!r}", location=loc)

        level: int | None = None
        level_text = get("level").strip()
        if "level" in cols and level_text:
            try:
                level = int(Decimal(level_text.lstrip(".")))
            except (InvalidOperation, ValueError) as exc:
                raise BOMParseError(f"层级不是整数: {level_text!r}", location=loc) from exc

        attributes: dict[str, str] = {}
        attr_text = get("attributes").strip()
        if attr_text:
            parsed = _parse_attributes(attr_text)
            attributes.update(parsed if parsed is not None else {"spec": attr_text})
        drawing = get("drawing").strip()
        if drawing:
            attributes["drawing"] = drawing
        for idx, name in header.extra.items():
            if idx < len(values) and values[idx] is not None and str(values[idx]).strip():
                attributes[name] = _cell_text(values[idx]).strip()

        rows.append(
            _Row(
                excel_row=excel_row,
                line_no=get("line_no").strip() or None,
                level=level,
                parent=get("parent").strip() or None,
                part_no=raw_part_no.strip(),
                part_desc=raw_desc.strip(),
                qty=qty,
                uom=get("uom").strip().upper() or DEFAULT_UOM,
                source=get("source").strip() or None,
                indent=indent,
                attributes=attributes,
            )
        )
    if header.currency is None and "currency" in cols:
        idx = cols["currency"]
        for row_values in ws.iter_rows(min_row=header.row_idx + 1, values_only=True):
            if idx < len(row_values) and row_values[idx]:
                header.currency = str(row_values[idx]).strip().upper()
                break
    return rows


def _pick_strategy(header: _Header, rows: Sequence[_Row]) -> LevelStrategy:
    if "level" in header.columns and any(r.level is not None for r in rows):
        return "level_column"
    if any(r.indent for r in rows):
        return "indent"
    if "parent" in header.columns and any(r.parent for r in rows):
        return "parent_column"
    return "flat"


def _levels_from_indent(rows: Sequence[_Row]) -> list[int]:
    ranks = {width: rank for rank, width in enumerate(sorted({r.indent for r in rows}))}
    return [ranks[r.indent] for r in rows]


def _levels_from_column(rows: Sequence[_Row]) -> list[int]:
    levels: list[int] = []
    for row in rows:
        if row.level is None:
            raise BOMParseError("level 列存在空值", location=f"R{row.excel_row}")
        levels.append(row.level)
    return levels


def _build_tree(
    rows: Sequence[_Row],
    strategy: LevelStrategy,
    rfq_id: str,
    project_name: str,
    currency: str | None,
) -> BOMTree:
    line_nos = _assign_line_nos(rows)
    if strategy in ("level_column", "indent"):
        levels = (
            _levels_from_column(rows) if strategy == "level_column" else _levels_from_indent(rows)
        )
        parents, levels, synth_root = _parents_from_levels(rows, levels, line_nos)
    elif strategy == "parent_column":
        parents, levels, synth_root = _parents_from_parent_column(rows, line_nos)
    else:
        levels = [1] * len(rows)
        parents = [None] * len(rows)
        synth_root = True

    root_line_no = _synth_root_line_no(line_nos) if synth_root else None
    nodes: list[BOMLine] = []
    if root_line_no is not None:
        nodes.append(
            BOMLine(
                line_no=root_line_no,
                level=0,
                parent_line_no=None,
                part_no=rfq_id,
                part_desc=project_name,
                material_class=classify_material(project_name),
                qty_per=Decimal(1),
                uom=DEFAULT_UOM,
                source_hint=None,
                auto_grade_required=False,
                single_source_risk=None,
                attributes={"synthesized": "true"},
            )
        )
    for row, line_no, level, parent in zip(rows, line_nos, levels, parents, strict=True):
        if level > 0 and parent is None:
            parent = root_line_no
        text = " ".join([row.part_desc, row.part_no, *row.attributes.values()])
        try:
            nodes.append(
                BOMLine(
                    line_no=line_no,
                    level=level,
                    parent_line_no=parent,
                    part_no=row.part_no,
                    part_desc=row.part_desc,
                    material_class=classify_material(f"{row.part_desc} {row.part_no}"),
                    qty_per=row.qty,
                    uom=row.uom,
                    source_hint=row.source,
                    auto_grade_required=is_auto_grade(f"{text} {row.source or ''}"),
                    single_source_risk=single_source_risk(row.source),
                    attributes=row.attributes,
                )
            )
        except ValidationError as exc:
            raise BOMParseError(f"BOM 行校验失败: {exc}", location=f"R{row.excel_row}") from exc

    root = next(n for n in nodes if n.level == 0)
    try:
        return BOMTree(root=root, nodes=nodes, currency_hint=currency, total_lines=len(nodes))
    except ValidationError as exc:
        raise BOMParseError(f"BOMTree 校验失败: {exc}") from exc


def _assign_line_nos(rows: Sequence[_Row]) -> list[str]:
    if all(r.line_no for r in rows):
        line_nos = [r.line_no or "" for r in rows]
    else:
        line_nos = [str(i) for i in range(1, len(rows) + 1)]
    seen: set[str] = set()
    for row, line_no in zip(rows, line_nos, strict=True):
        if line_no in seen:
            raise BOMParseError(f"行号重复: {line_no}", location=f"R{row.excel_row}")
        seen.add(line_no)
    return line_nos


def _synth_root_line_no(existing: Iterable[str]) -> str:
    taken = set(existing)
    return next(c for c in ("0", "ROOT", "ROOT-0") if c not in taken)


def _parents_from_levels(
    rows: Sequence[_Row], levels: list[int], line_nos: list[str]
) -> tuple[list[str | None], list[int], bool]:
    base = min(levels)
    levels = [lv - base for lv in levels]
    single_top = levels.count(0) == 1 and levels[0] == 0
    if not single_top:
        levels = [lv + 1 for lv in levels]  # 多个顶层或首行非顶层 → 合成 root
    parents: list[str | None] = []
    stack: list[str] = []  # stack[i] = 当前 level=i(+偏移) 的最近行号
    offset = 0 if single_top else 1
    for row, line_no, level in zip(rows, line_nos, levels, strict=True):
        depth = level - offset
        if depth > len(stack):
            raise BOMParseError(
                f"层级跳变：level={level} 缺少上一级父件", location=f"R{row.excel_row}"
            )
        del stack[depth:]
        parents.append(stack[-1] if stack else None)
        stack.append(line_no)
    return parents, levels, not single_top


def _parents_from_parent_column(
    rows: Sequence[_Row], line_nos: list[str]
) -> tuple[list[str | None], list[int], bool]:
    by_part: dict[str, str] = {}
    for row, line_no in zip(rows, line_nos, strict=True):
        by_part.setdefault(row.part_no, line_no)
    parents: list[str | None] = []
    for row in rows:
        if row.parent is None:
            parents.append(None)
            continue
        parent_line = by_part.get(row.parent)
        if parent_line is None:
            raise BOMParseError(f"父件号不存在: {row.parent}", location=f"R{row.excel_row}")
        parents.append(parent_line)

    tops = [i for i, p in enumerate(parents) if p is None]
    synth_root = len(tops) != 1
    parent_of = dict(zip(line_nos, parents, strict=True))
    levels: list[int] = []
    for row, line_no in zip(rows, line_nos, strict=True):
        depth, cursor, visited = 0, parent_of[line_no], {line_no}
        while cursor is not None:
            if cursor in visited:
                raise BOMParseError("父件关系成环", location=f"R{row.excel_row}")
            visited.add(cursor)
            depth += 1
            cursor = parent_of[cursor]
        levels.append(depth + (1 if synth_root else 0))
    return parents, levels, synth_root
